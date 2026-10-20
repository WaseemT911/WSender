import time
import pywhatkit as pwk
import sys
import os
import glob
import subprocess
import configparser
from rich import print

if os.name != "nt":
    import termios
    import tty


version = "0.5 beta"
verDate = "(21 Oct, 2026)"


FROZEN = getattr(sys, 'frozen', False)
if FROZEN:
    os.chdir(os.path.dirname(sys.executable))
else:
    os.chdir(os.path.dirname(os.path.abspath(__file__)))


if sys.stdout is not None:
    sys.stdout.reconfigure(encoding="utf-8")


DEFAULT_CONFIG = """[settings]
load_data = data.xlsx
wait_time = 10
tab_close = True
image_folder = images/
page_size = 25
"""


def _clean(value):
    return value.strip().strip('"').strip("'")


def config_file():
    if not os.path.exists('config.ini'):
        with open('config.ini', 'w', encoding='utf-8') as f:
            f.write(DEFAULT_CONFIG)
    config = configparser.ConfigParser(inline_comment_prefixes=('#', ';'))
    config.read('config.ini', encoding='utf-8')
    return {
        'LOAD_DATA': _clean(config.get('settings', 'load_data', fallback="data.xlsx")),
        'WAIT_TIME': config.getint('settings', 'wait_time', fallback=10),
        'TAB_CLOSE': config.getboolean('settings', 'tab_close', fallback=True),
        'IMAGE_FOLDER': _clean(config.get('settings', 'image_folder', fallback="images/")),
        'PAGE_SIZE': config.getint('settings', 'page_size', fallback=25),
    }


_cfg = config_file()
LOAD_DATA = _cfg['LOAD_DATA']
WAIT_TIME = _cfg['WAIT_TIME']
TAB_CLOSE = _cfg['TAB_CLOSE']
IMAGE_FOLDER = _cfg['IMAGE_FOLDER']
PAGE_SIZE = _cfg['PAGE_SIZE']


def clearConsole():
    os.system('cls' if os.name == 'nt' else 'clear')


def get_message_from_txt():
    if not os.path.exists('message.txt'):
        with open('message.txt', 'w', encoding='utf-8') as file:
            file.write("Hello!")
    with open('message.txt', 'r', encoding='utf-8') as file:
        message = file.read()
    return message


def get_images():
    if not os.path.isdir(IMAGE_FOLDER):
        os.makedirs(IMAGE_FOLDER, exist_ok=True)
        return None
    paths = [p for p in glob.glob(os.path.join(IMAGE_FOLDER, "*")) if os.path.isfile(p)]
    return sorted(paths)[-1] if paths else None


def send_whatsapp_msg(phone_number, message):
    try:
        pwk.sendwhatmsg_instantly(phone_number, message, wait_time=WAIT_TIME, tab_close=TAB_CLOSE)
        print(f"[blue]Message sent to {phone_number}!\n")
    except Exception as e:
        print(f"[red]Error: {e}")


def send_whatsapp_msg_with_image(phone_number, image, message):
    if sys.platform.startswith("linux"):
        subprocess.run(["copyq", "--start-server"], capture_output=True)

    try:
        pwk.sendwhats_image(phone_number, image, message, wait_time=WAIT_TIME, tab_close=TAB_CLOSE)
        print(f"[blue]Message sent to {phone_number}!\n")
    except Exception as e:
        print(f"[red]Error: {e}")


def get_user_choice(options):
    if os.name == "nt":
        import msvcrt
        while True:
            ch = msvcrt.getch().decode(errors="ignore")
            if ch in options:
                return ch
    else:
        fd = sys.stdin.fileno()
        old = termios.tcgetattr(fd)
        try:
            tty.setraw(fd)
            while True:
                ch = sys.stdin.read(1)
                if ch in options:
                    return ch
        finally:
            termios.tcsetattr(fd, termios.TCSADRAIN, old)


def input_with_esc(prompt):
    print(prompt, end="", flush=True)
    buffer = ""

    if os.name == "nt":
        import msvcrt
        while True:
            ch = msvcrt.getch()
            if ch in (b"\r", b"\n"):
                print()
                return buffer
            if ch == b"\x1b":
                print()
                return None
            if ch in (b"\x08", b"\x7f"):
                if buffer:
                    buffer = buffer[:-1]
                    sys.stdout.write("\b \b")
                    sys.stdout.flush()
                continue
            if ch in (b"\xe0", b"\x00"):
                msvcrt.getch()
                continue
            char = ch.decode(errors="ignore")
            if char:
                buffer += char
                sys.stdout.write(char)
                sys.stdout.flush()
    else:
        fd = sys.stdin.fileno()
        old = termios.tcgetattr(fd)
        try:
            tty.setraw(fd)
            while True:
                ch = sys.stdin.read(1)
                if ch in ("\r", "\n"):
                    sys.stdout.write("\r\n")
                    return buffer
                if ch == "\x1b":
                    sys.stdout.write("\r\n")
                    return None
                if ch == "\x03":
                    raise KeyboardInterrupt
                if ch in ("\x7f", "\x08"):
                    if buffer:
                        buffer = buffer[:-1]
                        sys.stdout.write("\b \b")
                        sys.stdout.flush()
                    continue
                buffer += ch
                sys.stdout.write(ch)
                sys.stdout.flush()
        finally:
            termios.tcsetattr(fd, termios.TCSADRAIN, old)


def get_customer_info(number_of_customers_input):
    import openpyxl

    global customerName
    global customerNum
    global rows_to_move

    if not os.path.exists(LOAD_DATA):
        wb = openpyxl.Workbook()
        wb.active.title = 'data'
        wb.create_sheet('finished')
        wb.save(LOAD_DATA)

    wb = openpyxl.load_workbook(LOAD_DATA)
    ws = wb['data']

    if 'finished' in wb.sheetnames: pass
    else: wb.create_sheet('finished')

    target = wb['finished']

    rows_to_move = []
    for row_index in range(1, ws.max_row + 1):
        if ws.cell(row=row_index, column=2).value is not None:
            rows_to_move.append(row_index)
            if len(rows_to_move) == number_of_customers_input:
                break

    customers_name = []
    customers_number = []
    for row_index in rows_to_move:
        customers_name.append((ws.cell(row=row_index, column=1).value))
        customers_number.append((ws.cell(row=row_index, column=2).value))
        target.append([ws.cell(row=row_index, column=col).value for col in range(1, ws.max_column + 1)])

    for row_index in reversed(rows_to_move):
        ws.delete_rows(row_index)

    wb.save(LOAD_DATA)

    def get_customer_name(): return customers_name
    def get_customer_number(): return customers_number

    return get_customer_name, get_customer_number


def about():
    clearConsole()
    print("""
                          ▓▓█
      ████              ██████           ▓█▓▓
     ▓█████▓           ██████          █████
       █████▓        ▓█████         ██████
         █████     ██████         ▓██████
          █████ ▓█████████▓      █████▓
           ██████████  ▓█████   █████
            ██████▓        ▓█▓█████▓
              ▓██▓           ████▓

          P   r   o   j   e   c   t   s
    ------------------------------------------
            """)
    print(f"Creator: WProjects by WaseemT911\nVersion: {version} {verDate}")
    input("\nPress Enter to go back...")
    main()


def call_customer_info_func():
    user_input = input_with_esc("Enter number of customers: ")
    if user_input is None:
        return None
    get_customer_name, get_customer_number = get_customer_info(int(user_input))
    return get_customer_name(), get_customer_number()


def main():
    clearConsole()
    username = os.getlogin()

    message = get_message_from_txt()
    image = get_images()

    print(f"Welcome {username} ;)\n")
    print("""
[1] Send message
[2] Send message with image
[3] Add / Edit Contacts
[4] About
[0] Exit
        """)

    print("\n[cyan]What do you want to do?")

    choice = get_user_choice("12340")

    if choice == '1':
        clearConsole()
        customer_info = call_customer_info_func()
        if customer_info is None:
            return main()
        names, numbers = customer_info
        for customerName, customerNum in zip(names, numbers):
            print(f"{customerName}: +962{customerNum}")
            send_whatsapp_msg(f"+962{customerNum}", message)
        input("\n[green][✓] Finished. Press Enter to go back...")
        main()

    elif choice == '2':
        clearConsole()
        if image is None:
            print(f"[red]No image found in '{IMAGE_FOLDER}'. Put an image there and try again.")
            input("\nPress Enter to go back...")
            return main()
        customer_info = call_customer_info_func()
        if customer_info is None:
            return main()
        names, numbers = customer_info
        for customerName, customerNum in zip(names, numbers):
            print(f"{customerName}: +962{customerNum}")
            send_whatsapp_msg_with_image(f"+962{customerNum}", image, message)
        input("\n[green][✓] Finished. Press Enter to go back...")
        main()

    elif choice == '3':
        if FROZEN:
            import gui
            app = gui.AddWindow()
            app.mainloop()
        else:
            script = os.path.join(os.path.dirname(os.path.abspath(__file__)), "gui.py")
            subprocess.run([sys.executable, script], capture_output=True)
        main()

    elif choice == '4': about()

    elif choice == '0': sys.exit(0)


if __name__ == "__main__":
    if "--gui" in sys.argv:
        import gui
        app = gui.AddWindow()
        app.mainloop()
    else:
        main()
