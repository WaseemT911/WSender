import os
import sys
import threading
import configparser
import customtkinter as ctk
import openpyxl

if getattr(sys, 'frozen', False):
    os.chdir(os.path.dirname(sys.executable))
else:
    os.chdir(os.path.dirname(os.path.abspath(__file__)))

try:
    import arabic_reshaper
    from bidi.algorithm import (
        get_empty_storage, get_base_level, get_embedding_levels,
        explicit_embed_and_overrides, resolve_weak_types, resolve_neutral_types,
        resolve_implicit_levels, reorder_resolved_levels, apply_mirroring,
        get_display,
    )
    _reshaper = arabic_reshaper.ArabicReshaper({'support_ligatures': True})
    _entry_reshaper = arabic_reshaper.ArabicReshaper({'support_ligatures': False})
    _have_bidi = True

    def ar(text):
        return get_display(_reshaper.reshape(str(text)))

    def _shape_1to1(raw):
        return _entry_reshaper.reshape(raw)

    def _bidi_map(shaped):
        if not shaped:
            return "", [], [], 0
        storage = get_empty_storage()
        base_level = get_base_level(shaped)
        storage["base_level"] = base_level
        storage["base_dir"] = ("L", "R")[base_level]
        get_embedding_levels(shaped, storage)
        for i, ch in enumerate(storage["chars"]):
            ch["orig_idx"] = i
        explicit_embed_and_overrides(storage, False)
        resolve_weak_types(storage, False)
        resolve_neutral_types(storage, False)
        resolve_implicit_levels(storage, False)
        levels = [ch["level"] for ch in storage["chars"]]
        reorder_resolved_levels(storage, False)
        apply_mirroring(storage, False)
        display = "".join(c["ch"] for c in storage["chars"])
        d2l = [c["orig_idx"] for c in storage["chars"]]
        return display, d2l, levels, base_level

    def _display_from_raw(raw):
        return _bidi_map(_shape_1to1(raw))

    def _disp_to_raw(dc, d2l, levels, n_disp):
        if n_disp == 0:
            return 0
        if dc >= n_disp:
            p = d2l[n_disp - 1]
            return p + 1 if levels[p] % 2 == 0 else p
        if dc <= 0:
            p = d2l[0]
            return p if levels[p] % 2 == 0 else p + 1
        p = d2l[dc - 1]
        return p + 1 if levels[p] % 2 == 0 else p

    def _raw_to_disp(rc, d2l, levels, n_raw):
        if n_raw == 0:
            return 0
        l2d = [0] * n_raw
        for di, si in enumerate(d2l):
            l2d[si] = di
        if rc >= n_raw:
            last = n_raw - 1
            return l2d[last] + 1 if levels[last] % 2 == 0 else l2d[last]
        return l2d[rc] if levels[rc] % 2 == 0 else l2d[rc] + 1

    def _backspace_raw_index(dc, d2l, levels, n_disp):
        """Raw index of the char Backspace deletes at display pos dc, or None."""
        if n_disp == 0:
            return None
        if dc > 0:
            if levels[d2l[dc - 1]] % 2 == 0:
                return d2l[dc - 1]
            if dc < n_disp:
                return d2l[dc]
            return None
        if levels[d2l[0]] % 2 == 1:
            return d2l[0]
        return None

    def _delete_raw_index(dc, d2l, levels, n_disp):
        """Raw index of the char Delete deletes at display pos dc, or None."""
        if n_disp == 0:
            return None
        if dc > 0 and levels[d2l[dc - 1]] % 2 == 1:
            return d2l[dc - 1]
        if dc < n_disp:
            return d2l[dc]
        if dc > 0:
            return d2l[dc - 1]
        return None

except ImportError:
    _have_bidi = False

    def ar(text):
        return text

    def _display_from_raw(raw):
        n = len(raw)
        return raw, list(range(n)), [0] * n, 0

    def _disp_to_raw(dc, d2l, levels, n_disp):
        return min(dc, n_disp)

    def _raw_to_disp(rc, d2l, levels, n_raw):
        return min(rc, n_raw)

    def _backspace_raw_index(dc, d2l, levels, n_disp):
        return dc - 1 if dc > 0 else None

    def _delete_raw_index(dc, d2l, levels, n_disp):
        return dc if dc < n_disp else None

_DEFAULT_CONFIG = """[settings]
load_data = data.xlsx
wait_time = 10
tab_close = True
image_folder = images/
page_size = 25
"""


def _load_config():
    if not os.path.exists('config.ini'):
        with open('config.ini', 'w', encoding='utf-8') as f:
            f.write(_DEFAULT_CONFIG)
    config = configparser.ConfigParser(inline_comment_prefixes=('#', ';'))
    config.read('config.ini', encoding='utf-8')
    return {
        'EXCEL_FILE': config.get('settings', 'load_data', fallback="data.xlsx").strip().strip('"').strip("'"),
        'PAGE_SIZE': config.getint('settings', 'page_size', fallback=25),
    }


_cfg = _load_config()
EXCEL_FILE = _cfg['EXCEL_FILE']
PAGE_SIZE = _cfg['PAGE_SIZE']
SHEET_NAME = 'data'
FINISHED_SHEET = 'finished'
ARABIC_FONT = "Tajawal"
global_font = "Quicksand Medium"


ctk.set_appearance_mode("dark")
ctk.set_default_color_theme("blue")


def get_workbook():
    if not os.path.exists(EXCEL_FILE):
        wb = openpyxl.Workbook()
        wb.active.title = SHEET_NAME
        wb.create_sheet(FINISHED_SHEET)
        wb.save(EXCEL_FILE)
    wb = openpyxl.load_workbook(EXCEL_FILE)
    if SHEET_NAME not in wb.sheetnames:
        wb.create_sheet(SHEET_NAME)
    if FINISHED_SHEET not in wb.sheetnames:
        wb.create_sheet(FINISHED_SHEET)
    return wb, wb[SHEET_NAME]


class ArabicEntry(ctk.CTkEntry):
    """CTkEntry that displays Arabic text properly shaped and reordered.

    Tk on Linux does not shape/reorder Arabic, so the typed text is kept raw
    internally and only the display is reshaped/reordered. get() returns the
    raw text. Display positions are mapped to raw positions so Backspace,
    Delete, selection and mid-text typing all work correctly.
    """

    def __init__(self, master, **kwargs):
        kwargs.setdefault("font", (ARABIC_FONT, 16))
        self._raw = ""
        self._updating = False
        super().__init__(master, **kwargs)
        self._refresh()
        self.bind("<KeyPress>", self._on_key, add="+")
        self.bind("<<Paste>>", self._on_paste, add="+")
        self.bind("<Button-2>", self._on_paste, add="+")

    def _selection(self):
        try:
            if self.select_present():
                return self.index("sel.first"), self.index("sel.last")
        except Exception:
            pass
        return None

    def _raw_indices_for_display_range(self, ds, de, d2l):
        idxs = sorted({d2l[i] for i in range(ds, de) if 0 <= i < len(d2l)})
        return idxs

    def _delete_raw_range(self, idxs, rc):
        if not idxs:
            return rc
        for i in reversed(idxs):
            self._raw = self._raw[:i] + self._raw[i + 1:]
        n_before = sum(1 for i in idxs if i < rc)
        return rc - n_before

    def _on_key(self, event):
        if self._updating:
            return
        if event.state & 0x4:
            if event.keysym == "v":
                return self._on_paste()
            if event.keysym == "c":
                self._copy_selection()
                return "break"
            if event.keysym == "x":
                sel = self._selection()
                if sel:
                    self._copy_selection()
                    display, d2l, levels, _ = _display_from_raw(self._raw)
                    idxs = self._raw_indices_for_display_range(sel[0], sel[1], d2l)
                    rc = _disp_to_raw(sel[0], d2l, levels, len(display))
                    self._delete_raw_range(idxs, rc)
                    self._refresh()
                return "break"
            if event.keysym == "a":
                return
            if event.keysym in ("z", "y"):
                return "break"
            return

        display, d2l, levels, _ = _display_from_raw(self._raw)
        n_disp = len(display)
        sel = self._selection()
        if sel:
            ds, de = sel
            dc = ds
        else:
            try:
                dc = self.index("insert")
            except Exception:
                dc = n_disp
            ds = de = None

        rc = _disp_to_raw(dc, d2l, levels, n_disp)

        if event.keysym == "BackSpace":
            if sel:
                idxs = self._raw_indices_for_display_range(ds, de, d2l)
                rc = self._delete_raw_range(idxs, rc)
            else:
                ri = _backspace_raw_index(dc, d2l, levels, n_disp)
                if ri is None:
                    return "break"
                self._raw = self._raw[:ri] + self._raw[ri + 1:]
                rc = ri
        elif event.keysym == "Delete":
            if sel:
                idxs = self._raw_indices_for_display_range(ds, de, d2l)
                rc = self._delete_raw_range(idxs, rc)
            else:
                ri = _delete_raw_index(dc, d2l, levels, n_disp)
                if ri is None:
                    return "break"
                self._raw = self._raw[:ri] + self._raw[ri + 1:]
                rc = ri
        elif event.keysym == "Escape":
            self._raw = ""
            rc = 0
        elif len(event.char) == 1 and event.char.isprintable():
            if sel:
                idxs = self._raw_indices_for_display_range(ds, de, d2l)
                rc = self._delete_raw_range(idxs, rc)
            self._raw = self._raw[:rc] + event.char + self._raw[rc:]
            rc += 1
        else:
            return

        self._refresh(rc)
        return "break"

    def _copy_selection(self):
        sel = self._selection()
        if not sel:
            return
        display, d2l, _, _ = _display_from_raw(self._raw)
        idxs = self._raw_indices_for_display_range(sel[0], sel[1], d2l)
        text = "".join(self._raw[i] for i in idxs)
        try:
            self.clipboard_clear()
            self.clipboard_append(text)
        except Exception:
            pass

    def _on_paste(self, event=None):
        try:
            text = self.clipboard_get()
        except Exception:
            return "break"
        text = text.replace("\n", " ").replace("\r", "")
        if not text:
            return "break"
        display, d2l, levels, _ = _display_from_raw(self._raw)
        sel = self._selection()
        if sel:
            idxs = self._raw_indices_for_display_range(sel[0], sel[1], d2l)
            rc = self._delete_raw_range(idxs, _disp_to_raw(sel[0], d2l, levels, len(display)))
        else:
            try:
                dc = self.index("insert")
            except Exception:
                dc = len(display)
            rc = _disp_to_raw(dc, d2l, levels, len(display))
        self._raw = self._raw[:rc] + text + self._raw[rc:]
        self._refresh(rc + len(text))
        return "break"

    def _refresh(self, raw_cursor=None):
        if raw_cursor is None:
            raw_cursor = len(self._raw)
        raw_cursor = max(0, min(raw_cursor, len(self._raw)))
        display, d2l, levels, _ = _display_from_raw(self._raw)
        self._updating = True
        try:
            self.delete(0, "end")
            if display:
                self.insert(0, display)
            self.icursor(_raw_to_disp(raw_cursor, d2l, levels, len(self._raw)))
        finally:
            self._updating = False

    def get(self):
        return self._raw

    def set(self, text):
        self._raw = str(text)
        self._refresh()

    def clear(self):
        self._raw = ""
        self._refresh()


class AddWindow(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("Add Contact")
        self.geometry("450x200")
        self.resizable(False, False)

        name_row = ctk.CTkFrame(self, fg_color="transparent")
        name_row.pack(fill="x", padx=20, pady=(20, 20))
        ctk.CTkLabel(name_row, text="Customer Name:", font=(global_font, 16), width=120, anchor="e").pack(side="left")
        self.name_entry = ArabicEntry(name_row, width=220, height=35, justify="center")
        self.name_entry.pack(side="left", padx=(4, 0), fill="x", expand=True)

        phone_row = ctk.CTkFrame(self, fg_color="transparent")
        phone_row.pack(fill="x", padx=20, pady=5)
        ctk.CTkLabel(phone_row, text="Phone Number:", font=(global_font, 16), width=120, anchor="e").pack(side="left")
        self.phone_entry = ArabicEntry(phone_row, width=220, height=35, font=(global_font, 16), justify="center")
        self.phone_entry.pack(side="left", padx=(10, 0), fill="x", expand=True)

        self.status_label = ctk.CTkLabel(self, text="", text_color="green", font=(global_font, 14))
        self.status_label.pack(pady=0)

        self._preview = None

        btn_frame = ctk.CTkFrame(self, fg_color="transparent")
        btn_frame.pack(pady=10)
        ctk.CTkButton(btn_frame, text="Add", command=self.add_contact,
                        fg_color="#444444", font=(global_font, 16), width=100, height=35).pack(side="left", padx=10)
        ctk.CTkButton(btn_frame, text="Preview Data", command=self.open_preview,
                        fg_color="#444444", font=(global_font, 16), width=150, height=35).pack(side="left", padx=10)

    def add_contact(self):
        name = self.name_entry.get().strip()
        phone = self.phone_entry.get().strip()

        if not name or not phone:
            self.status_label.configure(text="Please fill both fields", text_color="red")
            return

        wb, ws = get_workbook()
        ws.append([name, phone])
        wb.save(EXCEL_FILE)

        self.name_entry.clear()
        self.phone_entry.clear()
        self.status_label.configure(text=f"Added", text_color="green")

    def open_preview(self):
        if self._preview is None or not self._preview.winfo_exists():
            self._preview = PreviewWindow(self)
        else:
            self._preview.focus()
            self._preview.load_data()


class EditWindow(ctk.CTkToplevel):
    def __init__(self, preview, row_idx, name, phone):
        super().__init__(preview)
        self.preview = preview
        self.row_idx = row_idx
        self.title("Edit Contact")
        self.geometry("450x200")
        self.resizable(False, False)

        name_row = ctk.CTkFrame(self, fg_color="transparent")
        name_row.pack(fill="x", padx=20, pady=(20, 20))
        ctk.CTkLabel(name_row, text="Customer Name:", font=(global_font, 16), width=120, anchor="e").pack(side="left")
        self.name_entry = ArabicEntry(name_row, width=220, height=35, justify="center")
        self.name_entry.set(name)
        self.name_entry.pack(side="left", padx=(4, 0), fill="x", expand=True)

        phone_row = ctk.CTkFrame(self, fg_color="transparent")
        phone_row.pack(fill="x", padx=20, pady=5)
        ctk.CTkLabel(phone_row, text="Phone Number:", font=(global_font, 16), width=120, anchor="e").pack(side="left")
        self.phone_entry = ArabicEntry(phone_row, width=220, height=35, font=(global_font, 16), justify="center")
        self.phone_entry.set(phone)
        self.phone_entry.pack(side="left", padx=(10, 0), fill="x", expand=True)

        self.status_label = ctk.CTkLabel(self, text="", text_color="green", font=(global_font, 14))
        self.status_label.pack(pady=0)

        btn_frame = ctk.CTkFrame(self, fg_color="transparent")
        btn_frame.pack(pady=10)
        ctk.CTkButton(btn_frame, text="Save", command=self.save,
                      font=(global_font, 16), width=100, height=35).pack(side="left", padx=10)
        ctk.CTkButton(btn_frame, text="Cancel", command=self.destroy,
                      fg_color="#444444", font=(global_font, 16), width=100, height=35).pack(side="left", padx=10)

    def save(self):
        name = self.name_entry.get().strip()
        phone = self.phone_entry.get().strip()

        if not name or not phone:
            self.status_label.configure(text="Please fill both fields", text_color="red")
            return

        wb, ws = get_workbook()
        ws.cell(row=self.row_idx, column=1, value=name)
        ws.cell(row=self.row_idx, column=2, value=phone)
        wb.save(EXCEL_FILE)

        self.preview.load_data()
        self.destroy()




class PreviewWindow(ctk.CTkToplevel):
    def __init__(self, master):
        super().__init__(master)
        self.title("Preview Data")
        self.geometry("580x480")
        self.minsize(600, 500)
        self.resizable(False, False)
        self._view = "data"
        self._page = 0
        self._contacts = []
        self._loading = False

        top_row = ctk.CTkFrame(self, fg_color="transparent")
        top_row.pack(fill="x", padx=20, pady=(12, 0))
        ctk.CTkLabel(top_row, text=f"Contacts in {os.path.basename(EXCEL_FILE)}",
                     font=(global_font, 18)).pack(side="left")
        ctk.CTkButton(top_row, text="Refresh", command=self.load_data,
                      font=(global_font, 14), width=90, height=30).pack(side="right")

        self.view_switch = ctk.CTkSegmentedButton(
            self, values=["Data", "Finished"],
            command=self._on_view_change, font=(global_font, 14))
        self.view_switch.set("Data")
        self.view_switch.pack(pady=(0, 4))

        self.count_label = ctk.CTkLabel(self, text="", font=(global_font, 13),
                                        text_color="gray")
        self.count_label.pack(pady=(0, 8))

        header = ctk.CTkFrame(self, fg_color="transparent")
        header.pack(fill="x", padx=20)
        header.grid_columnconfigure(1, weight=1)
        ctk.CTkLabel(header, text="#", font=(global_font, 14, "bold"),
                     width=30, anchor="w").grid(row=0, column=0, sticky="ew")
        ctk.CTkLabel(header, text="Customer Name", font=(global_font, 14, "bold"),
                     anchor="w").grid(row=0, column=1, sticky="ew")
        ctk.CTkLabel(header, text="Phone Number", font=(global_font, 14, "bold"),
                     anchor="e").grid(row=0, column=2, sticky="ew", padx=(10, 8))
        self._action_header = ctk.CTkLabel(header, text="", font=(global_font, 14, "bold"),
                                           width=80)
        self._action_header.grid(row=0, column=3, sticky="e", padx=(0, 8))

        self.scroll = ctk.CTkScrollableFrame(self, width=530, height=240)
        self.scroll.pack(fill="both", expand=True, padx=20, pady=(4, 6))
        self.scroll.grid_columnconfigure(0, weight=1)

        self.loading_bar = ctk.CTkProgressBar(self, mode="indeterminate", height=6)
        self.loading_label = ctk.CTkLabel(self, text="", font=(global_font, 12),
                                          text_color="gray")

        self.page_frame = ctk.CTkFrame(self, fg_color="transparent")
        self.prev_btn = ctk.CTkButton(self.page_frame, text="◀ Prev", width=80, height=28,
                                      font=(global_font, 13), fg_color="#444444",
                                      command=self._prev_page)
        self.prev_btn.pack(side="left", padx=6)
        ctk.CTkLabel(self.page_frame, text="Page", font=(global_font, 13)).pack(side="left", padx=(8, 4))
        self.page_entry = ctk.CTkEntry(self.page_frame, width=50, height=28,
                                       font=(global_font, 13), justify="center")
        self.page_entry.pack(side="left", padx=2)
        self.page_entry.bind("<Return>", lambda e: self._goto_page())
        self.page_total_label = ctk.CTkLabel(self.page_frame, text="", font=(global_font, 13))
        self.page_total_label.pack(side="left", padx=(6, 8))
        self.next_btn = ctk.CTkButton(self.page_frame, text="Next ▶", width=80, height=28,
                                      font=(global_font, 13), fg_color="#444444",
                                      command=self._next_page)
        self.next_btn.pack(side="left", padx=6)

        self.load_data()

    def _on_view_change(self, value):
        self._view = value.lower()
        self._page = 0
        self.load_data()

    def _prev_page(self):
        if self._page > 0:
            self._page -= 1
            self._render_page()

    def _next_page(self):
        if (self._page + 1) * PAGE_SIZE < len(self._contacts):
            self._page += 1
            self._render_page()

    def _goto_page(self):
        try:
            page = int(self.page_entry.get().strip())
        except ValueError:
            self.page_entry.delete(0, "end")
            self.page_entry.insert(0, str(self._page + 1))
            return
        total_pages = self._total_pages()
        page = max(1, min(page, total_pages))
        if page - 1 != self._page:
            self._page = page - 1
            self._render_page()
        else:
            self.page_entry.delete(0, "end")
            self.page_entry.insert(0, str(page))

    def _total_pages(self):
        return max(1, (len(self._contacts) + PAGE_SIZE - 1) // PAGE_SIZE)

    def _get_contacts(self, sheet_name):
        try:
            wb = openpyxl.load_workbook(EXCEL_FILE, read_only=True)
        except Exception:
            return []
        if sheet_name not in wb.sheetnames:
            wb.close()
            return []
        ws = wb[sheet_name]
        contacts = []
        for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_col=2, values_only=True), start=1):
            name, phone = (row + (None, None))[:2] if row else (None, None)
            if name is None and phone is None:
                continue
            contacts.append((row_idx, str(name or ""), str(phone or "")))
        wb.close()
        return contacts

    def _show_loading(self, show):
        self._loading = show
        if show:
            self.scroll.pack_forget()
            self.page_frame.pack_forget()
            self.loading_label.configure(text="Loading contacts...")
            self.loading_label.pack(pady=(8, 2))
            self.loading_bar.pack(fill="x", padx=40)
            self.loading_bar.start()
        else:
            self.loading_bar.stop()
            self.loading_bar.pack_forget()
            self.loading_label.pack_forget()
            self.scroll.pack(fill="both", expand=True, padx=20, pady=(4, 6))

    def load_data(self):
        if self._loading:
            return
        self._show_loading(True)
        sheet = FINISHED_SHEET if self._view == "finished" else SHEET_NAME
        threading.Thread(target=self._load_worker, args=(sheet,), daemon=True).start()

    def _load_worker(self, sheet):
        contacts = self._get_contacts(sheet)
        self.after(0, lambda: self._load_done(contacts, sheet))

    def _load_done(self, contacts, sheet):
        if not self.winfo_exists():
            return
        self._contacts = contacts
        max_page = self._total_pages() - 1
        if self._page > max_page:
            self._page = max(0, max_page)
        self._show_loading(False)
        self.count_label.configure(
            text=f"{len(contacts)} contact(s) in '{sheet}'")
        self._action_header.configure(text="Actions")
        self._render_page()

    def _render_page(self):
        for widget in self.scroll.winfo_children():
            widget.destroy()

        is_finished = self._view == "finished"
        total = len(self._contacts)
        total_pages = self._total_pages()
        start = self._page * PAGE_SIZE
        page_rows = self._contacts[start:start + PAGE_SIZE]

        if total == 0:
            ctk.CTkLabel(self.scroll, text="No contacts yet",
                         font=(global_font, 15), text_color="gray").grid(
                row=0, column=0, pady=20)
            self.page_frame.pack_forget()
            return

        for i, (row_idx, name, phone) in enumerate(page_rows):
            line_no = start + i + 1
            row_frame = ctk.CTkFrame(
                self.scroll, height=34,
                fg_color=("gray90", "gray17") if i % 2 == 0 else "transparent")
            row_frame.grid(row=i, column=0, sticky="ew", pady=1)
            row_frame.grid_columnconfigure(1, weight=1)
            row_frame.grid_columnconfigure(2, weight=1)

            ctk.CTkLabel(row_frame, text=str(line_no), font=(global_font, 13),
                         text_color="gray", width=30, anchor="w").grid(
                row=0, column=0, sticky="ew", padx=(12, 0), pady=6)
            ctk.CTkLabel(row_frame, text=ar(name), font=(ARABIC_FONT, 15),
                         anchor="w").grid(row=0, column=1, sticky="ew", padx=8, pady=6)
            ctk.CTkLabel(row_frame, text=phone, font=(global_font, 15),
                         anchor="e").grid(row=0, column=2, sticky="e", padx=(10, 8), pady=6)

            if is_finished:
                ctk.CTkButton(row_frame, text="Restore", width=70, height=24,
                              font=(global_font, 12), fg_color="#2a7d2a",
                              command=lambda r=row_idx: self.restore_row(r)
                              ).grid(row=0, column=3, padx=(0, 8), pady=6)
            else:
                ctk.CTkButton(row_frame, text="Edit", width=50, height=24,
                              font=(global_font, 12), fg_color="#444444",
                              command=lambda r=row_idx, n=name, p=phone: EditWindow(self, r, n, p)
                              ).grid(row=0, column=3, padx=(4, 0), pady=6)
                ctk.CTkButton(row_frame, text="Delete", width=56, height=24,
                              font=(global_font, 12), fg_color="#a13d3d",
                              command=lambda r=row_idx: self.delete_row(r)
                              ).grid(row=0, column=4, padx=(5, 8), pady=6)

        self.page_total_label.configure(text=f"/ {total_pages}")
        self.page_entry.delete(0, "end")
        self.page_entry.insert(0, str(self._page + 1))
        self.prev_btn.configure(state="normal" if self._page > 0 else "disabled")
        self.next_btn.configure(
            state="normal" if self._page + 1 < total_pages else "disabled")
        self.page_frame.pack(pady=(0, 4))

    def delete_row(self, row_idx):
        wb = openpyxl.load_workbook(EXCEL_FILE)
        if SHEET_NAME not in wb.sheetnames:
            return
        ws = wb[SHEET_NAME]
        values = [ws.cell(row=row_idx, column=c).value for c in (1, 2)]
        if values[0] is None and values[1] is None:
            return
        ws.delete_rows(row_idx)
        wb.save(EXCEL_FILE)
        self.load_data()

    def restore_row(self, row_idx):
        wb = openpyxl.load_workbook(EXCEL_FILE)
        if FINISHED_SHEET not in wb.sheetnames:
            return
        ws_f = wb[FINISHED_SHEET]
        values = [ws_f.cell(row=row_idx, column=c).value for c in (1, 2)]
        if values[0] is None and values[1] is None:
            return

        if SHEET_NAME not in wb.sheetnames:
            wb.create_sheet(SHEET_NAME)
        wb[SHEET_NAME].append(values)
        ws_f.delete_rows(row_idx)
        wb.save(EXCEL_FILE)
        self.load_data()


if __name__ == "__main__":
    app = AddWindow()
    app.mainloop()
